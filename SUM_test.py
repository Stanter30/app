import os
from PIL import Image
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import pandas as pd
from pyzbar import pyzbar

image_folder = "output_A4"
excel_file = "report.xlsx"

image_files = [f for f in os.listdir(image_folder) if f.lower().endswith('.png')]
image_files.sort()

sums = []
sheets = []
sheets_n =[]
fd_dict = {}
sum_dict = {}

for img_file in image_files:

    img = Image.open(image_folder + '/' +img_file)
    qr_codes = pyzbar.decode(img)

    sum = 0
    fds = []
    for qr_code in qr_codes:
        text = qr_code.data.decode('utf-8')
        parts = text.split('s=')[1].split('&')
        price_value = float(parts[0])
        operation_type = int(parts[-1][-1])
        if operation_type == 2:
            price_value = -price_value
        sum += price_value

        fd = int(parts[2][2:])
        fds.append(fd)

    fd_dict[int(img_file[:3])] = fds
    sum_dict[int(img_file[:3])] = sum

# Запись листа в excel отчет
wb = load_workbook(excel_file)
sheet = wb.active
df = pd.read_excel(excel_file, header=None)
fd_list = df.iloc[3:, 1].tolist()

reverse_map = {}
for key, values in fd_dict.items():
    for val in values:
        reverse_map[val] = key

new_list = [reverse_map.get(num, None) for num in fd_list]

for index, value in enumerate(new_list):
    current_row = 4 + index
    cell_coordinate = f"{"L"}{current_row}"
    sheet[cell_coordinate] = value

#sheet.unmerge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
#sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
sheet.merge_cells(start_row=2, start_column=12, end_row=3, end_column=12)
sheet['L2'].value = "Лист А4"
sheet['L2'].font = Font(bold=True, size=11, name='Arial')
sheet['L2'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb.save(excel_file)